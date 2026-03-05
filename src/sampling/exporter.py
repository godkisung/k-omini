"""검수 결과 Excel 내보내기 모듈.

``openpyxl``을 사용하여 납품업체 피드백용 Excel 파일을 생성합니다.

시트 구성:
    - 요약: 배치명, 검수 일자, 유형별 오류율 통계
    - 검수결과: 파일별 상세 결과 (오류 여부, 유형, 메모)
"""

from __future__ import annotations

import os
from datetime import datetime
from io import BytesIO
from typing import TYPE_CHECKING

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

if TYPE_CHECKING:
    from .cache_manager import SamplingCache
    from .sampler import DOC_TYPE_NAMES


# ── 색상 팔레트 ──────────────────────────────────────────────────────────────
_CLR_HEADER_BG = "2F4F7F"   # 헤더 배경 (진파랑)
_CLR_HEADER_FG = "FFFFFF"   # 헤더 글자 (흰색)
_CLR_OK_BG = "D6FFD6"       # 정상 행 배경 (연초록)
_CLR_ERR_BG = "FFD6D6"      # 오류 행 배경 (연빨강)
_CLR_SECTION = "E8F0FE"     # 섹션 구분 배경 (연파랑)


def _header_style() -> tuple[Font, PatternFill, Alignment]:
    """공통 헤더 스타일을 반환합니다."""
    font = Font(bold=True, color=_CLR_HEADER_FG, size=11)
    fill = PatternFill("solid", fgColor=_CLR_HEADER_BG)
    align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    return font, fill, align


def _apply_header(ws, row: int, col: int, value: str) -> None:
    """셀에 헤더 스타일을 적용합니다."""
    cell = ws.cell(row=row, column=col, value=value)
    font, fill, align = _header_style()
    cell.font = font
    cell.fill = fill
    cell.alignment = align


def _set_col_widths(ws, widths: list[int]) -> None:
    """열 너비를 일괄 설정합니다."""
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


class ReviewExporter:
    """검수 결과를 Excel 파일로 내보냅니다.

    Args:
        cache: 검수 완료된 SamplingCache 인스턴스.
    """

    def __init__(self, cache: "SamplingCache") -> None:
        self.cache = cache

    # ── 공개 API ─────────────────────────────────────────────────────────────

    def to_bytes(self) -> bytes:
        """Excel 파일을 bytes 형태로 반환합니다 (Streamlit 다운로드용).

        Returns:
            xlsx 바이너리.
        """
        wb = self._build_workbook()
        buf = BytesIO()
        wb.save(buf)
        return buf.getvalue()

    def save(self, output_path: str) -> str:
        """Excel 파일을 지정 경로에 저장합니다.

        Args:
            output_path: 저장할 파일 경로 (.xlsx).

        Returns:
            저장된 파일의 절대 경로.
        """
        wb = self._build_workbook()
        os.makedirs(os.path.dirname(output_path), exist_ok=True)
        wb.save(output_path)
        return os.path.abspath(output_path)

    # ── 내부 구현 ─────────────────────────────────────────────────────────────

    def _build_workbook(self) -> openpyxl.Workbook:
        """워크북 전체를 구성합니다."""
        wb = openpyxl.Workbook()
        wb.remove(wb.active)  # 기본 시트 제거

        self._build_summary_sheet(wb)
        self._build_detail_sheet(wb)

        return wb

    # ── 요약 시트 ─────────────────────────────────────────────────────────────

    def _build_summary_sheet(self, wb: openpyxl.Workbook) -> None:
        """'요약' 시트를 생성합니다."""
        ws = wb.create_sheet("요약")
        cache = self.cache
        reviewed, total = cache.progress()

        # ── 기본 정보 ──────────────────────────────────────────────────────
        info_rows = [
            ("납품 배치", cache.batch_name),
            ("검수 일자", cache.date),
            ("검수 완료일", datetime.now().strftime("%Y-%m-%d %H:%M")),
            ("전체 샘플 수", total),
            ("검수 완료 수", reviewed),
            ("샘플링 비율", f"{cache._data.get('sample_rate', 0) * 100:.0f}%"),
            ("난수 시드", cache._data.get("seed", "")),
        ]

        ws.column_dimensions["A"].width = 20
        ws.column_dimensions["B"].width = 40

        for r, (label, value) in enumerate(info_rows, start=1):
            ws.cell(row=r, column=1, value=label).font = Font(bold=True)
            ws.cell(row=r, column=2, value=value)

        # ── 유형별 오류율 표 ───────────────────────────────────────────────
        start_row = len(info_rows) + 2
        headers = ["문서 유형", "유형명", "전체 페이지", "샘플 수", "오류 수", "오류율(%)"]
        for c, h in enumerate(headers, start=1):
            _apply_header(ws, start_row, c, h)

        # 유형별 오류 집계
        type_errors: dict[str, int] = {}
        for fname, result in cache.review_results.items():
            if not result.get("reviewed", False):
                continue
            doc_type = fname[:2].upper()
            if result.get("is_error", False):
                type_errors[doc_type] = type_errors.get(doc_type, 0) + 1

        row = start_row + 1
        total_errors = 0
        for doc_type, info in sorted(cache.type_breakdown.items()):
            errors = type_errors.get(doc_type, 0)
            total_errors += errors
            sampled = info.get("sampled", 0)
            error_rate = round(errors / sampled * 100, 1) if sampled > 0 else 0.0

            ws.cell(row=row, column=1, value=doc_type)
            ws.cell(row=row, column=2, value=info.get("name", doc_type))
            ws.cell(row=row, column=3, value=info.get("total", 0))
            ws.cell(row=row, column=4, value=sampled)
            ws.cell(row=row, column=5, value=errors)
            ws.cell(row=row, column=6, value=error_rate)

            # 오류율에 따른 행 색상
            fill_color = _CLR_ERR_BG if error_rate > 0 else _CLR_OK_BG
            fill = PatternFill("solid", fgColor=fill_color)
            for c in range(1, 7):
                ws.cell(row=row, column=c).fill = fill
            row += 1

        # 합계 행
        ws.cell(row=row, column=1, value="합계").font = Font(bold=True)
        ws.cell(row=row, column=4, value=total).font = Font(bold=True)
        ws.cell(row=row, column=5, value=total_errors).font = Font(bold=True)
        overall_rate = round(total_errors / total * 100, 1) if total > 0 else 0.0
        ws.cell(row=row, column=6, value=overall_rate).font = Font(bold=True)

        _set_col_widths(ws, [10, 24, 14, 10, 10, 12])

    # ── 상세 검수결과 시트 ────────────────────────────────────────────────────

    def _build_detail_sheet(self, wb: openpyxl.Workbook) -> None:
        """'검수결과' 시트를 생성합니다."""
        ws = wb.create_sheet("검수결과")

        headers = [
            "번호", "파일명", "문서 유형", "유형명",
            "오류 여부", "오류 유형", "검수자 메모", "검수 완료",
        ]
        for c, h in enumerate(headers, start=1):
            _apply_header(ws, 1, c, h)

        from .sampler import DOC_TYPE_NAMES as _TYPES

        for i, fname in enumerate(self.cache.sampled_files, start=1):
            result = self.cache.get_result(fname)
            doc_type = fname[:2].upper()
            is_error = result.get("is_error", False)
            reviewed = result.get("reviewed", False)

            row_data = [
                i,
                fname,
                doc_type,
                _TYPES.get(doc_type, doc_type),
                "오류" if is_error else "정상",
                ", ".join(result.get("error_types", [])),
                result.get("memo", ""),
                "✅" if reviewed else "⬜",
            ]

            fill_color = (
                _CLR_ERR_BG if is_error
                else (_CLR_OK_BG if reviewed else "FFFFFF")
            )
            fill = PatternFill("solid", fgColor=fill_color)

            for c, val in enumerate(row_data, start=1):
                cell = ws.cell(row=i + 1, column=c, value=val)
                cell.fill = fill
                cell.alignment = Alignment(vertical="center", wrap_text=(c == 2))

        _set_col_widths(ws, [6, 60, 10, 18, 10, 30, 30, 10])
        ws.row_dimensions[1].height = 28
